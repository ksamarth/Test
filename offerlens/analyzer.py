"""Claude-powered offer analysis.

One structured call does both stages: extract the contract into a normalized
schema, then produce red flags, negotiation points, and attorney questions.
PDFs go to Claude natively as document blocks; DOCX/TXT are sent as text.
"""

import base64
import io
import json
import os
from pathlib import Path

import anthropic

from schemas import OfferAnalysis

MODEL = "claude-opus-4-8"

SYSTEM_PROMPT = """You are the analysis engine for OfferLens, a tool that helps dental and \
medical students and new graduates understand their employment offers. Your reader has strong \
clinical training and essentially no business, legal, or financial background. They may be \
comparing several offers and deciding what to negotiate and what to ask a real attorney about.

You are an educator and preparation aid, not a lawyer, accountant, or financial advisor:
- Explain what clauses mean and what attorneys typically flag in similar contracts. Never give \
directives like "do not sign this" or "you should accept" — instead frame guidance as what to \
verify, what to ask, and what is commonly negotiated.
- Judge the contract against what is typical for new-graduate healthcare employment agreements \
(dental associate contracts, physician employment agreements, DSO agreements, hospital employment).
- Treat missing terms as findings: if a contract is silent on tail coverage, patient record \
ownership, or how production is calculated, that silence is itself worth flagging.
- Be honest about strengths. A tool that calls everything a red flag teaches the reader nothing.
- Write in plain English. Define any term of art the first time you use it (e.g. "tail coverage — \
insurance that covers claims filed after you leave").
- Extract numbers exactly as written. Use null for anything the document does not state; never \
guess a value.

Classic traps to check for in dental/medical offers: collections-based pay presented as \
production-based; lab fees deducted before the production split; daily guarantees that expire \
quickly or are repayable; non-competes measured from every office of a multi-site DSO; \
claims-made malpractice with tail coverage owed by the provider; sign-on bonus clawbacks with \
long service commitments; vague for-cause termination triggers; 1099 status for what is \
functionally employment; auto-renewal terms with narrow exit windows; the employer keeping all \
patient records and goodwill."""

_DEMO = os.environ.get("OFFERLENS_DEMO", "").lower() in ("1", "true", "yes")


def _document_blocks(filename: str, data: bytes) -> list:
    """Turn an uploaded file into Claude content blocks."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return [
            {
                "type": "document",
                "source": {
                    "type": "base64",
                    "media_type": "application/pdf",
                    "data": base64.standard_b64encode(data).decode("utf-8"),
                },
            }
        ]
    if suffix == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        text = "\n".join(p.text for p in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text for cell in row.cells)
        return [{"type": "text", "text": text}]
    # Fall back to treating the bytes as UTF-8 text (.txt and similar)
    return [{"type": "text", "text": data.decode("utf-8", errors="replace")}]


def analyze_offer(filename: str, data: bytes) -> OfferAnalysis:
    if _DEMO:
        return _demo_analysis()

    client = anthropic.Anthropic()
    content = _document_blocks(filename, data)
    content.append(
        {
            "type": "text",
            "text": "Analyze this employment offer for a healthcare new graduate. "
            "Extract every stated term into the schema, then assess it.",
        }
    )

    response = client.messages.parse(
        model=MODEL,
        max_tokens=16000,
        thinking={"type": "adaptive"},
        system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": content}],
        output_format=OfferAnalysis,
    )
    if response.stop_reason == "refusal":
        raise RuntimeError("The model declined to analyze this document.")
    return response.parsed_output


def _demo_analysis() -> OfferAnalysis:
    """Canned analysis of the bundled sample offer, for running without an API key."""
    demo_path = Path(__file__).parent / "sample_offers" / "demo_analysis.json"
    return OfferAnalysis.model_validate(json.loads(demo_path.read_text()))
